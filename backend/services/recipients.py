from pathlib import Path

from openpyxl import load_workbook


RECIPIENT_FILE = Path("part2/recipients.xlsx")


def load_recipients(
    file_path: str | Path = RECIPIENT_FILE,
) -> list[dict[str, str]]:
    """Load expense report recipients from an Excel workbook."""

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook["Recipients"]

    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        return []

    headers = [str(value).strip().lower() if value else "" for value in rows[0]]

    recipients = []

    for row in rows[1:]:
        values = dict(zip(headers, row))

        name = str(values.get("name") or "").strip()
        email = str(values.get("email") or "").strip()

        if name and email:
            recipients.append({
                "name": name,
                "email": email,
            })

    return recipients